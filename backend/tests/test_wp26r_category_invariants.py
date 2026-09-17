"""WP26R §1/§3 -- category-list invariants: the shared validator, and every
question create/edit/import/rename write path that must enforce it. Offline,
temp DB only; never touches the real ``backend/src/database/app.db``.
"""

from __future__ import annotations

import json

import pytest

from src.utils.category_order import CATEGORY_ORDER
from src.utils.category_rules import (
    CategoryValidationError,
    dedupe_and_validate,
    validate_categories,
)

CAT_A = CATEGORY_ORDER[0]
CAT_B = CATEGORY_ORDER[1]
CAT_C = CATEGORY_ORDER[2]


# --------------------------------------------------------------------------- #
# the shared validator itself
# --------------------------------------------------------------------------- #
def test_validate_categories_accepts_valid_ordered_list():
    assert validate_categories([CAT_A, CAT_B]) == [CAT_A, CAT_B]


def test_validate_categories_rejects_empty_list():
    with pytest.raises(CategoryValidationError):
        validate_categories([])


def test_validate_categories_rejects_non_list():
    with pytest.raises(CategoryValidationError):
        validate_categories(CAT_A)  # a bare string, not a list


def test_validate_categories_rejects_unknown_category():
    with pytest.raises(CategoryValidationError):
        validate_categories([CAT_A, "קטגוריה שלא קיימת"])


def test_validate_categories_rejects_duplicate():
    with pytest.raises(CategoryValidationError):
        validate_categories([CAT_A, CAT_A])


def test_validate_categories_rejects_empty_string_entry():
    with pytest.raises(CategoryValidationError):
        validate_categories([""])


def test_validate_categories_never_repairs_returns_fresh_list():
    src = [CAT_A, CAT_B]
    out = validate_categories(src)
    out.append(CAT_C)
    assert src == [CAT_A, CAT_B]  # caller's list untouched


def test_dedupe_and_validate_merges_duplicates_preserving_first_occurrence():
    assert dedupe_and_validate([CAT_A, CAT_B, CAT_A]) == [CAT_A, CAT_B]


def test_dedupe_and_validate_still_rejects_unknown_category():
    with pytest.raises(CategoryValidationError):
        dedupe_and_validate([CAT_A, "לא קנוני"])


def test_dedupe_and_validate_still_rejects_empty():
    with pytest.raises(CategoryValidationError):
        dedupe_and_validate([])


# --------------------------------------------------------------------------- #
# question create (Excel import) -- malformed row rejected, others still import
# --------------------------------------------------------------------------- #
def test_upload_excel_rejects_row_with_unknown_category_others_still_import(client):
    from openpyxl import Workbook
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.append(["נושא", "שאלה", "תשובה1", "תשובה2", "תשובה3", "תשובה4", "תשובה_נכונה"])
    ws.append([CAT_A, "שאלה תקינה", "a", "b", "c", "d", 1])
    ws.append(["קטגוריה לא קנונית", "שאלה עם קטגוריה שגויה", "a", "b", "c", "d", 1])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/upload-excel",
        data={"file": (buf, "q.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["success"] is True
    assert "1" in body["message"] or "הועלו 1" in body["message"]
    assert body["details"]["errors"]  # the bad row was reported, not silently dropped


def test_upload_excel_deduplicates_repeated_category_columns(client, app):
    """A row repeating the same category across נושא/נושא2 is a benign data
    -entry slip, not a malformed write -- the existing import dedup collapses
    it to a single canonical entry (still nonempty/unique/canonical) rather
    than rejecting the row."""
    from openpyxl import Workbook
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.append(["נושא", "שאלה", "תשובה1", "תשובה2", "תשובה3", "תשובה4", "תשובה_נכונה", "נושא2"])
    ws.append([CAT_A, "שאלה עם קטגוריה כפולה", "a", "b", "c", "d", 1, CAT_A])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = client.post(
        "/api/upload-excel",
        data={"file": (buf, "q.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    with app.app_context():
        from src.models.question import Question

        q = Question.query.order_by(Question.id.desc()).first()
        assert q.categories == [CAT_A]


# --------------------------------------------------------------------------- #
# question edit (PUT)
# --------------------------------------------------------------------------- #
def _create_one(client, categories=(CAT_A,)):
    resp = client.post(
        "/api/upload-excel",
        data={"file": _one_row_xlsx(categories)},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200
    from src.models.question import Question

    return Question.query.order_by(Question.id.desc()).first().id


def _one_row_xlsx(categories):
    from openpyxl import Workbook
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    headers = ["נושא", "שאלה", "תשובה1", "תשובה2", "תשובה3", "תשובה4", "תשובה_נכונה"]
    row = [categories[0], f"שאלה {categories}", "a", "b", "c", "d", 1]
    if len(categories) > 1:
        headers.append("נושא2")
        row.append(categories[1])
    ws.append(headers)
    ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return (buf, "q.xlsx")


def test_update_question_rejects_unknown_category(app, client):
    with app.app_context():
        qid = _create_one(client)
    resp = client.put(f"/api/questions/{qid}", json={"categories": ["לא קנוני"]})
    assert resp.status_code == 400
    with app.app_context():
        from src.models.question import Question

        assert Question.query.get(qid).categories == [CAT_A]  # untouched


def test_update_question_rejects_duplicate_categories(app, client):
    with app.app_context():
        qid = _create_one(client)
    resp = client.put(f"/api/questions/{qid}", json={"categories": [CAT_A, CAT_A]})
    assert resp.status_code == 400


def test_update_question_rejects_contradictory_singular_primary(app, client):
    with app.app_context():
        qid = _create_one(client, categories=(CAT_A, CAT_B))
    resp = client.put(
        f"/api/questions/{qid}",
        json={"category": CAT_B, "categories": [CAT_A, CAT_B]},  # contradicts categories[0]
    )
    assert resp.status_code == 400


def test_update_question_accepts_consistent_category_and_categories(app, client):
    with app.app_context():
        qid = _create_one(client, categories=(CAT_A, CAT_B))
    resp = client.put(
        f"/api/questions/{qid}",
        json={"category": CAT_A, "categories": [CAT_A, CAT_B]},
    )
    assert resp.status_code == 200
    with app.app_context():
        from src.models.question import Question

        q = Question.query.get(qid)
        assert q.category == CAT_A
        assert q.categories == [CAT_A, CAT_B]


# --------------------------------------------------------------------------- #
# secondary category add/remove
# --------------------------------------------------------------------------- #
def test_add_secondary_category_rejects_unknown_category(app, client):
    with app.app_context():
        qid = _create_one(client)
    resp = client.post(f"/api/questions/{qid}/add-category", json={"category": "לא קנוני"})
    assert resp.status_code == 400


def test_add_secondary_category_accepts_canonical_category(app, client):
    with app.app_context():
        qid = _create_one(client)
    resp = client.post(f"/api/questions/{qid}/add-category", json={"category": CAT_B})
    assert resp.status_code == 200
    assert resp.get_json()["categories"] == [CAT_A, CAT_B]


def test_remove_secondary_category_recomputes_primary(app, client):
    with app.app_context():
        qid = _create_one(client, categories=(CAT_A, CAT_B))
    resp = client.delete(f"/api/questions/{qid}/remove-category", json={"category": CAT_A})
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["categories"] == [CAT_B]
    assert body["category"] == CAT_B


# --------------------------------------------------------------------------- #
# category rename
# --------------------------------------------------------------------------- #
def test_rename_category_rejects_non_canonical_target(app, client):
    with app.app_context():
        _create_one(client, categories=(CAT_A,))
    resp = client.put(f"/api/categories/{CAT_A}", json={"new_name": "שם לא קנוני"})
    assert resp.status_code == 400
    with app.app_context():
        from src.models.question import Question

        assert Question.query.first().categories == [CAT_A]  # untouched


def test_rename_category_dedupes_on_merge_into_existing_category(app, client):
    with app.app_context():
        qid = _create_one(client, categories=(CAT_A, CAT_B))
    resp = client.put(f"/api/categories/{CAT_A}", json={"new_name": CAT_B})
    assert resp.status_code == 200
    with app.app_context():
        from src.models.question import Question

        q = Question.query.get(qid)
        assert q.categories == [CAT_B]  # deduplicated, not [CAT_B, CAT_B]
        assert q.category == CAT_B


def test_rename_category_preserves_order_and_recomputes_primary(app, client):
    with app.app_context():
        qid = _create_one(client, categories=(CAT_A, CAT_B))
    resp = client.put(f"/api/categories/{CAT_A}", json={"new_name": CAT_C})
    assert resp.status_code == 200
    with app.app_context():
        from src.models.question import Question

        q = Question.query.get(qid)
        assert q.categories == [CAT_C, CAT_B]
        assert q.category == CAT_C


# --------------------------------------------------------------------------- #
# browsing/counting stays consistent with the full list (not only primary)
# --------------------------------------------------------------------------- #
def test_categories_summary_counts_secondary_membership(app, client):
    with app.app_context():
        _create_one(client, categories=(CAT_A, CAT_B))
    resp = client.get("/api/categories")
    assert resp.status_code == 200
    names = {c["name"] for c in resp.get_json()}
    assert CAT_A in names and CAT_B in names


# --------------------------------------------------------------------------- #
# audit read-only guarantee: opening the DB in mode=ro cannot mutate it
# --------------------------------------------------------------------------- #
def test_readonly_audit_connection_cannot_write(app, tmp_path):
    """Mirrors the exact mechanism used for the real read-only category-model
    audit against ``app.db``: a ``sqlite3`` connection opened with
    ``mode=ro`` must refuse any write, proving the audit technique itself
    cannot mutate the database it inspects."""
    import sqlite3

    with app.app_context():
        from src.models.user import db

        db_path = str(db.engine.url.database)

    ro_conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    try:
        with pytest.raises(sqlite3.OperationalError):
            ro_conn.execute("INSERT INTO questions (id) VALUES (999999)")
    finally:
        ro_conn.close()
