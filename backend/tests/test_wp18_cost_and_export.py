"""WP18 §4/§5/§6 -- cumulative ceiling, terminal summary, export, legacy routes."""

from __future__ import annotations

import io
import json
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from tests._wp18_fakes import Dispatch, dispatch_factory

from src.jobs import service, store
from src.models.question import Question
from src.models.user import db


def _run(jobs_app, payload, *, provider_factory=None):
    with jobs_app.app_context():
        job = service.create_job(payload, seed=payload.get("_seed"))
    return service.run_job(job.job_id, provider_factory=provider_factory)


def test_cumulative_ceiling_stops_the_run_and_marks_the_slot(jobs_app, jobs_root, llm_ready):
    # each fake call costs ~ $0.34 conservative bound; $0.50 cap admits one slot
    done = _run(
        jobs_app,
        {"categories": {"היסטולוגיה": {"total": 4, "database": 1, "llm": 3}},
         "cost_ceiling_usd": "0.50", "_seed": 2},
        provider_factory=dispatch_factory(Dispatch()),
    )
    assert done.status == "cost_ceiling"
    llm = [s for s in done.slots if s.kind == "llm"]
    assert sum(1 for s in llm if s.status == "accepted") == 1
    assert any(s.status == "cost_ceiling" for s in llm)
    assert Decimal(done.accumulated_cost_usd) <= Decimal("0.50")
    # remaining slot never ran
    assert any(s.status == "queued" for s in llm)
    # the DB question is still there
    assert sum(1 for s in done.slots if s.kind == "database" and s.status == "accepted") == 1


def test_all_costs_counted_accepted_and_failed(jobs_app, jobs_root, llm_ready):
    from tests._wp18_fakes import RejectingProvider

    done = _run(
        jobs_app,
        {"categories": {"מבוא": {"total": 2, "database": 0, "llm": 2}}, "cost_ceiling_usd": "5.00"},
        provider_factory=dispatch_factory(Dispatch(factory=RejectingProvider)),
    )
    assert done.status == "failed"                      # nothing accepted
    # both failed attempts still cost something and were ledgered
    assert Decimal(done.accumulated_cost_usd) > 0
    assert len(done.cost_ledger) == 2
    assert all(e["status"] in ("question_rejected", "provider_output_failure") for e in done.cost_ledger)
    # cost sums exactly
    assert Decimal(done.accumulated_cost_usd) == sum(
        Decimal(e["total_cost_usd"]) for e in done.cost_ledger
    )


def test_terminal_summary_contents_and_no_secret(jobs_app, jobs_root, monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-SUPER-SECRET-do-not-persist")
    done = _run(
        jobs_app,
        {"categories": {"היסטולוגיה": {"total": 3, "database": 1, "llm": 2}},
         "cost_ceiling_usd": "5.00", "_seed": 4},
        provider_factory=dispatch_factory(Dispatch()),
    )
    s = done.terminal_summary
    assert s["job_id"] == done.job_id
    assert s["status"] == "completed"
    assert s["final_llm_cost_usd"] == done.accumulated_cost_usd
    assert s["cost_basis"] in ("calculated", "conservative_bound", "mixed")
    assert s["llm_accepted"] == 2 and s["llm_failed"] == 0
    assert "retries" in s and "pricing_warnings" in s and "pricing_verification" in s

    # the full persisted job carries no key value anywhere
    blob = (store.job_dir(done.job_id) / "job.json").read_text("utf-8")
    assert "sk-SUPER-SECRET" not in blob
    for p in store.job_dir(done.job_id).rglob("*"):
        if p.is_file():
            assert "sk-SUPER-SECRET" not in p.read_text("utf-8", errors="ignore"), p


def test_cost_ceiling_update_cannot_go_below_accumulated(jobs_app, jobs_root, llm_ready):
    done = _run(
        jobs_app,
        {"categories": {"מבוא": {"total": 1, "database": 0, "llm": 1}}, "cost_ceiling_usd": "5.00"},
        provider_factory=dispatch_factory(Dispatch()),
    )
    acc = Decimal(done.accumulated_cost_usd)
    assert acc > 0
    try:
        service.update_cost_ceiling(done.job_id, str(acc / 2))
        raised = False
    except service.JobConflict:
        raised = True
    assert raised
    # raising it is fine
    up = service.update_cost_ceiling(done.job_id, "9.99")
    assert up.cost_ceiling_usd == "9.99"


def test_export_only_accepted_llm_and_round_trips_through_upload(jobs_app, jobs_client, jobs_root, llm_ready):
    done = _run(
        jobs_app,
        {"categories": {"היסטולוגיה": {"total": 3, "database": 2, "llm": 1}},
         "cost_ceiling_usd": "5.00", "_seed": 6},
        provider_factory=dispatch_factory(Dispatch()),
    )
    data = service.export_llm_xlsx(done)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == tuple(service.EXPORT_HEADERS)
    assert len(rows) == 2  # header + exactly the 1 accepted llm question (no DB rows)

    with jobs_app.app_context():
        before = Question.query.count()
    resp = jobs_client.post(
        "/api/upload-excel",
        data={"file": (io.BytesIO(data), "generated.xlsx")},
        content_type="multipart/form-data",
    )
    assert resp.status_code == 200, resp.get_data(as_text=True)
    with jobs_app.app_context():
        after = Question.query.count()
        assert after == before + 1  # the export route did NOT insert; upload did
        newq = Question.query.order_by(Question.id.desc()).first()
        assert newq.category == "היסטולוגיה"
        assert newq.correct_answer_id in (1, 2, 3, 4)


def test_legacy_test_endpoints_still_work(jobs_client, jobs_app, jobs_root):
    cats = jobs_client.get("/api/test/categories")
    assert cats.status_code == 200
    names = [r["name"] for r in cats.get_json()]
    from src.utils.category_order import CATEGORY_ORDER
    assert names[: len(CATEGORY_ORDER)] == list(CATEGORY_ORDER)

    gen = jobs_client.post("/api/test/generate", json={"categories": {"היסטולוגיה": 2}})
    assert gen.status_code == 200
    body = gen.get_json()
    assert body["total_questions"] == 2
    assert all(q["number"] in (1, 2) for q in body["questions"])


def test_create_job_via_route_returns_202_and_polls_to_completion(jobs_client, jobs_app, jobs_root, llm_ready):
    jobs_app.config["EXAM_JOB_SYNC"] = True
    jobs_app.config["EXAM_JOB_PROVIDER_FACTORY"] = dispatch_factory(Dispatch())
    resp = jobs_client.post(
        "/api/exam-jobs",
        json={"categories": {"מבוא": {"total": 2, "database": 1, "llm": 1}}, "cost_ceiling_usd": "5.00"},
    )
    assert resp.status_code == 202
    body = resp.get_json()
    jid = body["job_id"]
    got = jobs_client.get(body["url"])
    assert got.status_code == 200
    view = got.get_json()
    # WP27: creation only selects DB questions (db_review); the planned LLM
    # slot needs an explicit continue-llm before the job reaches "completed".
    assert view["status"] == "queued"
    assert view["workflow_phase"] == "db_review"
    assert [q["origin"] for q in view["questions"]] == ["database"]

    cont = jobs_client.post(f"/api/exam-jobs/{jid}/continue-llm")
    assert cont.status_code == 202, cont.get_json()
    view = jobs_client.get(body["url"]).get_json()
    assert view["status"] == "completed"
    assert [q["origin"] for q in view["questions"]] == ["database", "llm"]
    assert view["terminal_summary"]["job_id"] == jid
