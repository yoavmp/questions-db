"""WP18R -- cross-source question replacement.

Both endpoints must accept **any** accepted current slot:

* ``replace-db``  -> DB->DB and LLM->DB
* ``replace-llm`` -> DB->LLM and LLM->LLM

with the semantic-history rule intact (only a displaced *LLM* question enters
``category_history``, and only after the replacement succeeds), stable
identity/number/category/order, correct per-origin state and DTO ``origin`` after
each transition, immutable historical cost ledger entries, export membership that
follows the *current* origin, and a full rollback on cross-source failure.

Temp DB, offline fake providers, sockets blocked (see ``conftest``). Zero
provider calls.
"""

from __future__ import annotations

from decimal import Decimal

import pytest

from tests._wp18_fakes import Dispatch, RejectingProvider, dispatch_factory

from src.jobs import service, store
from src.models.question import Question

CAT = "היסטולוגיה"


def _mixed(jobs_app, disp: Dispatch | None = None):
    """A completed job with 2 accepted DB + 2 accepted LLM questions in one
    category (the fixture seeds 6 DB rows there, so 4 remain for DB swaps)."""
    disp = disp or Dispatch()
    with jobs_app.app_context():
        job = service.create_job(
            {"categories": {CAT: {"total": 4, "database": 2, "llm": 2}},
             "cost_ceiling_usd": "5.00", "_seed": 11}
        )
    done = service.run_job(job.job_id, provider_factory=dispatch_factory(disp))
    assert done.status == "completed"
    return done


def _dto_for(job, instance_id: str) -> dict:
    return next(q for q in service.result_view(job)["questions"]
               if q["instance_id"] == instance_id)


# --------------------------------------------------------------------------- #
# the four-case replacement matrix
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("start_kind, action, new_kind, expect_history", [
    ("database", "replace_db",  "database", False),   # DB  -> DB   : no history
    ("database", "replace_llm", "llm",      False),   # DB  -> LLM  : no history
    ("llm",      "replace_db",  "database", True),    # LLM -> DB   : history after success
    ("llm",      "replace_llm", "llm",      True),    # LLM -> LLM  : history after success
])
def test_replacement_matrix(jobs_app, jobs_root, llm_ready, start_kind, action, new_kind, expect_history):
    done = _mixed(jobs_app)
    s = next(x for x in done.slots if x.kind == start_kind and x.status == "accepted")
    iid, num, cat, order = s.instance_id, s.number, s.category, s.order_in_category
    old_q = dict(s.question)

    if action == "replace_db":
        with jobs_app.app_context():
            up = service.replace_from_db(done.job_id, iid)
    else:
        up = service.replace_via_llm(done.job_id, iid, provider_factory=dispatch_factory(Dispatch()))

    ns = up.slot_by_instance(iid)
    # both endpoints accepted this starting origin
    assert ns is not None
    # stable identity / number / category / order
    assert (ns.instance_id, ns.number, ns.category, ns.order_in_category) == (iid, num, cat, order)
    # correct per-origin state after the transition
    assert ns.kind == new_kind
    if new_kind == "database":
        assert isinstance(ns.db_id, int)
        assert ns.was_repaired is False and ns.audit_ref is None      # no LLM metadata
    else:
        assert ns.db_id is None
    # DTO origin follows the current kind
    dto = _dto_for(up, iid)
    assert dto["origin"] == ("database" if new_kind == "database" else "llm")
    if new_kind == "database":
        assert "generation_meta" not in dto
    else:
        assert dto["generation_meta"]["outcome"] == "accepted"
    # the exact history matrix
    assert (old_q in up.category_history.get(CAT, [])) is expect_history


# --------------------------------------------------------------------------- #
# per-transition detail
# --------------------------------------------------------------------------- #
def test_db_to_db_costs_nothing_and_swaps_question(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    s = next(x for x in done.slots if x.kind == "database")
    iid, old_id = s.instance_id, s.db_id
    in_exam = {x.db_id for x in done.slots if x.kind == "database"}
    cost_before = done.accumulated_cost_usd

    with jobs_app.app_context():
        up = service.replace_from_db(done.job_id, iid)

    ns = up.slot_by_instance(iid)
    assert ns.db_id != old_id and ns.db_id not in in_exam
    assert up.accumulated_cost_usd == cost_before
    assert not any(e["kind"] == "replace_llm" for e in up.cost_ledger)


def test_db_to_llm_does_not_feed_displaced_db_to_generator(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    s = next(x for x in done.slots if x.kind == "database")
    iid = s.instance_id
    displaced = s.question["question"]
    cost_before = Decimal(done.accumulated_cost_usd)

    up = service.replace_via_llm(done.job_id, iid, provider_factory=dispatch_factory(Dispatch()))

    ns = up.slot_by_instance(iid)
    assert ns.kind == "llm" and ns.db_id is None
    assert ns.question["question"] != displaced
    # displaced DB question: not in history, never in a later generator context
    assert all(h["question"] != displaced for h in up.category_history.get(CAT, []))
    reloaded = store.load(done.job_id)
    other = next(x for x in reloaded.slots if x.kind == "llm" and x.instance_id != iid)
    prev = service._previous_for_slot(reloaded, other)
    assert all(p["question"] != displaced for p in prev)
    # LLM replacement uses + updates the cumulative ledger (incl. this attempt)
    assert Decimal(up.accumulated_cost_usd) > cost_before
    assert any(e["kind"] == "replace_llm" for e in up.cost_ledger)


def test_llm_to_db_appends_history_and_clears_llm_metadata(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")
    iid = s.instance_id
    old_q = dict(s.question)
    cost_before = done.accumulated_cost_usd
    ledger_len = len(done.cost_ledger)

    with jobs_app.app_context():
        up = service.replace_from_db(done.job_id, iid)

    ns = up.slot_by_instance(iid)
    assert ns.kind == "database" and isinstance(ns.db_id, int)
    assert ns.attempts == 0 and ns.retries == 0 and ns.was_repaired is False
    assert old_q in up.category_history[CAT]                 # displaced LLM -> history
    assert up.accumulated_cost_usd == cost_before           # DB adds no cost
    # historical ledger entries are immutable when the slot changes origin
    assert up.cost_ledger[:ledger_len] == done.cost_ledger[:ledger_len]


def test_llm_to_llm_appends_history_after_success(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")
    iid = s.instance_id
    old_q = dict(s.question)
    cost_before = Decimal(done.accumulated_cost_usd)

    up = service.replace_via_llm(done.job_id, iid, provider_factory=dispatch_factory(Dispatch()))

    ns = up.slot_by_instance(iid)
    assert ns.kind == "llm" and ns.db_id is None
    assert ns.question["question"] != old_q["question"]
    assert ns.question["number"] == old_q["number"]
    assert old_q in up.category_history[CAT]
    assert Decimal(up.accumulated_cost_usd) > cost_before


# --------------------------------------------------------------------------- #
# failure rollback
# --------------------------------------------------------------------------- #
def _export_rows(job) -> list:
    from io import BytesIO
    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(service.export_llm_xlsx(job))).active
    return list(ws.iter_rows(min_row=1, values_only=True))


def test_db_to_llm_failure_restores_slot_and_export(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    s = next(x for x in done.slots if x.kind == "database")
    iid = s.instance_id
    before = (s.kind, s.status, s.db_id, dict(s.question), s.attempts, s.retries,
              s.was_repaired, s.audit_ref)
    hist_before = list(done.category_history.get(CAT, []))
    export_before = _export_rows(store.load(done.job_id))
    cost_before = Decimal(done.accumulated_cost_usd)

    up = service.replace_via_llm(
        done.job_id, iid,
        provider_factory=dispatch_factory(Dispatch(factory=RejectingProvider)),
    )

    ns = up.slot_by_instance(iid)
    after = (ns.kind, ns.status, ns.db_id, dict(ns.question), ns.attempts, ns.retries,
             ns.was_repaired, ns.audit_ref)
    assert after == before                                  # slot state identical
    assert up.category_history.get(CAT, []) == hist_before   # no history mutation
    assert _export_rows(up) == export_before                 # projection unchanged
    # the failed attempt is still charged to the cumulative budget + ledger
    assert Decimal(up.accumulated_cost_usd) > cost_before
    assert any(e["kind"] == "replace_llm" for e in up.cost_ledger)


def test_llm_to_db_failure_no_alternative_keeps_llm_slot(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    llm_s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")
    in_exam = {x.db_id for x in done.slots if x.kind == "database"}

    with jobs_app.app_context():
        from src.models.user import db
        for q in Question.query.filter(Question.category == CAT).all():
            if q.id not in in_exam:
                db.session.delete(q)
        db.session.commit()
        with pytest.raises(service.JobConflict):
            service.replace_from_db(done.job_id, llm_s.instance_id)

    still = store.load(done.job_id).slot_by_instance(llm_s.instance_id)
    assert still.kind == "llm" and still.db_id is None
    assert still.question == llm_s.question
    assert store.load(done.job_id).category_history.get(CAT, []) == []


# --------------------------------------------------------------------------- #
# cross-cutting: context, ordering, export membership, DB isolation
# --------------------------------------------------------------------------- #
def test_displaced_llm_reaches_context_displaced_db_does_not(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    llm_s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")
    db_s = next(x for x in done.slots if x.kind == "database")
    displaced_llm = llm_s.question["question"]
    displaced_db = db_s.question["question"]

    with jobs_app.app_context():
        service.replace_from_db(done.job_id, llm_s.instance_id)          # LLM -> DB (llm -> history)
    service.replace_via_llm(done.job_id, db_s.instance_id,              # DB -> LLM (db NOT -> history)
                            provider_factory=dispatch_factory(Dispatch()))

    reloaded = store.load(done.job_id)
    target = next(x for x in reloaded.slots
                  if x.kind == "llm" and x.instance_id not in (llm_s.instance_id, db_s.instance_id))
    texts = [p["question"] for p in service._previous_for_slot(reloaded, target)]
    assert displaced_llm in texts
    assert displaced_db not in texts
    hist = reloaded.category_history[CAT]
    assert any(h["question"] == displaced_llm for h in hist)
    assert all(h["question"] != displaced_db for h in hist)


def test_numbers_category_and_order_stable_across_transitions(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    base = [(q["number"], q["category"], q["instance_id"])
            for q in service.result_view(done)["questions"]]
    assert [n for n, _, _ in base] == [1, 2, 3, 4]

    db_s = next(x for x in done.slots if x.kind == "database")
    llm_s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")
    service.replace_via_llm(done.job_id, db_s.instance_id, provider_factory=dispatch_factory(Dispatch()))
    with jobs_app.app_context():
        service.replace_from_db(done.job_id, llm_s.instance_id)

    after = [(q["number"], q["category"], q["instance_id"])
             for q in service.result_view(store.load(done.job_id))["questions"]]
    assert after == base                                    # number+category+order+identity intact


def test_export_membership_follows_current_origin(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    db_s = next(x for x in done.slots if x.kind == "database")
    llm_s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")

    # DB -> LLM: the new generated question joins the export
    up = service.replace_via_llm(done.job_id, db_s.instance_id,
                                 provider_factory=dispatch_factory(Dispatch()))
    new_llm_text = up.slot_by_instance(db_s.instance_id).question["question"]
    # LLM -> DB: that slot leaves the export; its displaced text stays only in history
    with jobs_app.app_context():
        up = service.replace_from_db(up.job_id, llm_s.instance_id)
    displaced_llm_text = llm_s.question["question"]

    from io import BytesIO
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(service.export_llm_xlsx(store.load(done.job_id))))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    exported_questions = [r[1] for r in rows]
    exported_instances = [s for s in store.load(done.job_id).slots if s.kind == "llm" and s.status == "accepted"]
    assert len(rows) == len(exported_instances)
    assert new_llm_text in exported_questions               # current origin llm -> exported
    assert displaced_llm_text not in exported_questions      # history-only -> not exported


def test_cross_source_ops_never_touch_appdb(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        before_count = Question.query.count()
        before_texts = {q.question for q in Question.query.all()}

    done = _mixed(jobs_app)
    db_s = next(x for x in done.slots if x.kind == "database")
    llm_s = next(x for x in done.slots if x.kind == "llm" and x.status == "accepted")
    gen_text = service.replace_via_llm(
        done.job_id, db_s.instance_id, provider_factory=dispatch_factory(Dispatch()),
    ).slot_by_instance(db_s.instance_id).question["question"]
    with jobs_app.app_context():
        service.replace_from_db(done.job_id, llm_s.instance_id)
        after_count = Question.query.count()
        after_texts = {q.question for q in Question.query.all()}

    assert after_count == before_count
    assert after_texts == before_texts
    assert gen_text not in after_texts


def test_both_endpoints_accept_both_origins_via_route(jobs_client, jobs_app, jobs_root, llm_ready):
    jobs_app.config["EXAM_JOB_SYNC"] = True
    jobs_app.config["EXAM_JOB_PROVIDER_FACTORY"] = dispatch_factory(Dispatch())
    r = jobs_client.post("/api/exam-jobs", json={
        "categories": {CAT: {"total": 4, "database": 2, "llm": 2}},
        "cost_ceiling_usd": "5.00", "_seed": 11,
    })
    assert r.status_code == 202
    jid = r.get_json()["job_id"]
    # WP27: creation only selects DB questions (db_review) -- the planned LLM
    # pair needs an explicit continue-llm before an origin="llm" question exists.
    cont = jobs_client.post(f"/api/exam-jobs/{jid}/continue-llm")
    assert cont.status_code == 202, cont.get_json()
    view = jobs_client.get(f"/api/exam-jobs/{jid}").get_json()
    db_q = next(q for q in view["questions"] if q["origin"] == "database")
    llm_q = next(q for q in view["questions"] if q["origin"] == "llm")

    # replace-llm accepts a DB question (DB -> LLM)
    resp = jobs_client.post(f"/api/exam-jobs/{jid}/questions/{db_q['instance_id']}/replace-llm")
    assert resp.status_code == 200, resp.get_data(as_text=True)
    # replace-db accepts an LLM question (LLM -> DB)
    resp = jobs_client.post(f"/api/exam-jobs/{jid}/questions/{llm_q['instance_id']}/replace-db")
    assert resp.status_code == 200, resp.get_data(as_text=True)

    final = jobs_client.get(f"/api/exam-jobs/{jid}").get_json()
    origin_by_iid = {q["instance_id"]: q["origin"] for q in final["questions"]}
    assert origin_by_iid[db_q["instance_id"]] == "llm"
    assert origin_by_iid[llm_q["instance_id"]] == "database"
    # numbers still contiguous 1..4, category order intact
    assert [q["number"] for q in final["questions"]] == [1, 2, 3, 4]
