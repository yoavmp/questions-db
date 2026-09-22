"""WP27 -- staged DB review before LLM generation.

Creating a mixed exam now selects DB questions only (``db_review`` /
``status == "queued"``); the originally planned LLM batch runs only after an
explicit ``continue_llm_generation`` / ``POST …/continue-llm`` call. Covers
every item in WP27's "Backend contract tests" list. Offline, deterministic:
temp DB, fake in-process providers (``tests._wp18_fakes``), sockets blocked.
Zero real provider calls.
"""

from __future__ import annotations

import io
import socket
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from tests._wp18_fakes import ApprovingProvider, Dispatch, dispatch_factory

from src.jobs import service, store

CAT = "היסטולוגיה"


@pytest.fixture(autouse=True)
def _no_network(monkeypatch):
    def _boom(*a, **k):  # pragma: no cover
        raise AssertionError("network access attempted during a WP27 test")

    monkeypatch.setattr(socket.socket, "connect", _boom)
    monkeypatch.setattr(socket.socket, "connect_ex", _boom)


# --------------------------------------------------------------------------- #
# 1. mixed A/B creation with no key and provider/readiness/audit traps
# --------------------------------------------------------------------------- #
def test_create_job_makes_zero_provider_readiness_or_audit_calls(jobs_app, jobs_root, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)

    def _boom_readiness(*a, **k):
        raise AssertionError("readiness_report() must not be called during create_job (WP27 §2)")

    def _boom_generate(*a, **k):
        raise AssertionError("generate_category_question() must not be called during create_job")

    def _boom_audit(*a, **k):
        raise AssertionError("no invocation audit dir may be created during create_job")

    monkeypatch.setattr(service, "readiness_report", _boom_readiness)
    monkeypatch.setattr(service, "generate_category_question", _boom_generate)
    monkeypatch.setattr(store, "invocation_audit_dir", _boom_audit)

    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 1, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 1,
        })
    assert job.status == "queued"
    assert service.workflow_phase(job) == "db_review"
    assert job.accumulated_cost_usd == "0"
    assert job.cost_ledger == []
    accepted = [s for s in job.slots if s.status == "accepted"]
    assert accepted and all(s.kind == "database" for s in accepted)
    pending = service.pending_llm_slots(job)
    assert len(pending) == 2 and all(s.status == "queued" and s.question is None for s in pending)


# --------------------------------------------------------------------------- #
# 2. B=0 completes immediately, no continuation action
# --------------------------------------------------------------------------- #
def test_b_zero_completes_immediately_no_continuation_action(jobs_app, jobs_root):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 2, "llm": 0}},
            "cost_ceiling_usd": "5.00", "_seed": 2,
        })
    assert job.status == "completed"
    assert service.workflow_phase(job) == "complete"
    assert job.terminal_summary is not None
    with pytest.raises(service.JobConflict):
        service.continue_llm_generation(job.job_id)


# --------------------------------------------------------------------------- #
# 3. A=0/B>0 creates a valid empty DB-review job
# --------------------------------------------------------------------------- #
def test_a_zero_b_positive_creates_empty_db_review_job(jobs_app, jobs_root):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 0, "llm": 2}},
            "cost_ceiling_usd": "5.00",
        })
    assert job.status == "queued"
    assert service.workflow_phase(job) == "db_review"
    view = service.result_view(job)
    assert view["questions"] == []
    assert view["pending_llm_total"] == 2
    assert view["accepted_count"] == 0


# --------------------------------------------------------------------------- #
# 4. Continue generates exactly the planned B items once
# --------------------------------------------------------------------------- #
def test_continue_generates_exactly_the_planned_b_items_once(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 4, "database": 2, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 3,
        })
    disp = Dispatch()
    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(disp))
    assert done.status == "completed"
    assert service.workflow_phase(done) == "complete"
    llm_slots = [s for s in done.slots if s.kind == "llm"]
    assert len(llm_slots) == 2 and all(s.status == "accepted" for s in llm_slots)
    assert sum(p.generation_calls for p in disp.by_context.values()) == 2

    # already complete -- a repeat Continue cannot reset/replay the batch
    with pytest.raises(service.JobConflict):
        service.continue_llm_generation(done.job_id, provider_factory=dispatch_factory(Dispatch()))


# --------------------------------------------------------------------------- #
# 5. double / concurrent Continue cannot duplicate slots, calls, cost, ledger
# --------------------------------------------------------------------------- #
def test_double_and_concurrent_continue_cannot_duplicate(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 4,
        })

    # concurrent: another operation holds the run lock
    assert service._RUN_LOCK.acquire(blocking=False)
    try:
        with pytest.raises(service.JobBusy):
            service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    finally:
        service._RUN_LOCK.release()
    untouched = store.load(job.job_id)
    assert untouched.status == "queued" and untouched.cost_ledger == []

    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    assert done.status == "completed"
    ledger_len = len(done.cost_ledger)
    cost = done.accumulated_cost_usd

    # double: a request once the job has already left db_review
    with pytest.raises(service.JobConflict):
        service.continue_llm_generation(done.job_id, provider_factory=dispatch_factory(Dispatch()))
    reloaded = store.load(done.job_id)
    assert len(reloaded.cost_ledger) == ledger_len
    assert reloaded.accumulated_cost_usd == cost
    assert sum(1 for s in reloaded.slots if s.kind == "llm" and s.status == "accepted") == 1


# --------------------------------------------------------------------------- #
# 6. readiness failure before a provider call leaves the job retryable, $0
# --------------------------------------------------------------------------- #
def test_readiness_failure_leaves_job_in_db_review_and_charges_nothing(jobs_app, jobs_root, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 5,
        })
    with pytest.raises(service.JobError) as exc:
        service.continue_llm_generation(job.job_id)  # no provider injected -> readiness gate fires
    assert "not ready" in str(exc.value)

    reloaded = store.load(job.job_id)
    assert reloaded.status == "queued"
    assert service.workflow_phase(reloaded) == "db_review"
    assert reloaded.accumulated_cost_usd == "0"
    assert reloaded.cost_ledger == []

    # safely retryable once the environment is fixed
    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    assert done.status == "completed"


# --------------------------------------------------------------------------- #
# 7. manual LLM replacement during db_review is charged, does not reduce B
# --------------------------------------------------------------------------- #
def test_manual_llm_replacement_during_db_review_does_not_reduce_planned_b(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 4, "database": 2, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 6,
        })
    assert service.workflow_phase(job) == "db_review"
    db_slot = next(s for s in job.slots if s.kind == "database")

    updated = service.replace_via_llm(job.job_id, db_slot.instance_id,
                                       provider_factory=dispatch_factory(Dispatch()))
    assert updated.status == "queued"
    assert service.workflow_phase(updated) == "db_review"  # unchanged phase after a manual op
    new_slot = updated.slot_by_instance(db_slot.instance_id)
    assert new_slot.kind == "llm" and new_slot.status == "accepted"
    assert Decimal(updated.accumulated_cost_usd) > 0

    plan = updated.plan_for(CAT)
    assert plan.database == 2 and plan.llm == 2  # original plan (provenance) untouched
    assert len(service.pending_llm_slots(updated)) == 2  # planned B slots untouched

    done = service.continue_llm_generation(updated.job_id, provider_factory=dispatch_factory(Dispatch()))
    assert done.status == "completed"
    accepted_llm = [s for s in done.slots if s.kind == "llm" and s.status == "accepted"]
    assert len(accepted_llm) == 3  # the 1 manual + the 2 originally planned


# --------------------------------------------------------------------------- #
# 8. both replacement directions during db_review preserve stable identity
# --------------------------------------------------------------------------- #
def test_both_replacement_directions_during_db_review_preserve_identity(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 2, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 7,
        })
    db_slot = next(s for s in job.slots if s.kind == "database")
    old_iid, old_num = db_slot.instance_id, db_slot.number

    with jobs_app.app_context():
        after_db = service.replace_from_db(job.job_id, old_iid)
    new_db_slot = after_db.slot_by_instance(old_iid)
    assert new_db_slot.instance_id == old_iid and new_db_slot.number == old_num
    assert new_db_slot.kind == "database"

    after_llm = service.replace_via_llm(job.job_id, old_iid, provider_factory=dispatch_factory(Dispatch()))
    new_llm_slot = after_llm.slot_by_instance(old_iid)
    assert new_llm_slot.instance_id == old_iid and new_llm_slot.number == old_num
    assert new_llm_slot.kind == "llm" and new_llm_slot.status == "accepted"


# --------------------------------------------------------------------------- #
# 9. displaced-LLM vs displaced-DB history rules hold during db_review
# --------------------------------------------------------------------------- #
def test_displaced_history_rules_hold_during_db_review(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 2, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 8,
        })
    db_slot = next(s for s in job.slots if s.kind == "database")

    # DB -> LLM: the displaced DB question never enters history
    updated = service.replace_via_llm(job.job_id, db_slot.instance_id,
                                       provider_factory=dispatch_factory(Dispatch()))
    assert updated.category_history[CAT] == []
    llm_slot = updated.slot_by_instance(db_slot.instance_id)

    # LLM -> DB: the displaced LLM question DOES enter history
    with jobs_app.app_context():
        updated2 = service.replace_from_db(updated.job_id, llm_slot.instance_id)
    assert len(updated2.category_history[CAT]) == 1


# --------------------------------------------------------------------------- #
# 10/11. semantic context: pending items excluded; history + earlier same-run
# accepted LLM questions included
# --------------------------------------------------------------------------- #
class CapturingDispatch(Dispatch):
    """Records exactly what ``request.previous_questions`` contains for each
    generation call, then delegates to a real ``ApprovingProvider`` so the
    call still succeeds normally -- same technique as
    ``test_wp25g_inverse_duplicate_guard.py``."""

    def __init__(self):
        super().__init__(factory=ApprovingProvider)
        self.seen_previous: list[list[dict]] = []

    def generate_candidates(self, request, *, prompts):
        self.seen_previous.append([
            {
                "number": p.number, "question": p.question,
                "answer1": p.answer1, "answer2": p.answer2,
                "answer3": p.answer3, "answer4": p.answer4,
                "correct_answer": p.correct_answer,
            }
            for p in request.previous_questions
        ])
        return super().generate_candidates(request, prompts=prompts)


def test_continuation_context_excludes_pending_includes_history_and_earlier_batch_items(
    jobs_app, jobs_root, llm_ready,
):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 1, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 9,
        })
    db_slot = next(s for s in job.slots if s.kind == "database")

    # seed one displaced-LLM history entry: DB -> LLM -> DB dance, all during
    # db_review, before Continue is ever claimed.
    after_llm = service.replace_via_llm(job.job_id, db_slot.instance_id,
                                         provider_factory=dispatch_factory(Dispatch()))
    llm_slot = after_llm.slot_by_instance(db_slot.instance_id)
    with jobs_app.app_context():
        after_db2 = service.replace_from_db(after_llm.job_id, llm_slot.instance_id)
    assert len(after_db2.category_history[CAT]) == 1
    assert service.workflow_phase(after_db2) == "db_review"
    assert len(service.pending_llm_slots(after_db2)) == 2  # both planned B slots still pending

    cap = CapturingDispatch()
    done = service.continue_llm_generation(after_db2.job_id, provider_factory=dispatch_factory(cap))
    assert done.status == "completed"
    assert len(cap.seen_previous) == 2

    # first planned LLM slot's context: the 1 current DB question + the 1
    # history entry -- exactly 2, nothing else (no pending-placeholder leak)
    first_ctx = cap.seen_previous[0]
    assert len(first_ctx) == 2
    numbers_seen = {q["number"] for q in first_ctx}
    assert after_db2.slot_by_instance(db_slot.instance_id).number in numbers_seen

    # second planned LLM slot's context additionally includes the FIRST
    # planned slot's own just-accepted question from this same continuation
    second_ctx = cap.seen_previous[1]
    assert len(second_ctx) == 3
    accepted_llm_slots = [s for s in done.slots if s.kind == "llm" and s.status == "accepted"]
    first_generated_number = min(s.number for s in accepted_llm_slots)
    assert first_generated_number in {q["number"] for q in second_ctx}


# --------------------------------------------------------------------------- #
# 12. manual pre-Continue spend and automatic continuation share ONE ceiling
# --------------------------------------------------------------------------- #
def test_manual_and_continuation_spend_share_one_cost_ceiling(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 1, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 10,
        })
    db_slot = next(s for s in job.slots if s.kind == "database")
    updated = service.replace_via_llm(job.job_id, db_slot.instance_id,
                                       provider_factory=dispatch_factory(Dispatch()))
    manual_spend = Decimal(updated.accumulated_cost_usd)
    assert manual_spend > 0

    done = service.continue_llm_generation(updated.job_id, provider_factory=dispatch_factory(Dispatch()))
    assert Decimal(done.accumulated_cost_usd) > manual_spend  # continuation adds to the SAME total
    assert done.cost_ceiling_usd == "5.00"  # not a fresh budget


def test_continuation_respects_a_ceiling_already_exhausted_manually(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 1, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 11,
        })
    db_slot = next(s for s in job.slots if s.kind == "database")
    updated = service.replace_via_llm(job.job_id, db_slot.instance_id,
                                       provider_factory=dispatch_factory(Dispatch()))
    manual_spend = Decimal(updated.accumulated_cost_usd)
    assert manual_spend > 0

    # tighten the ceiling to just barely above what's already spent -- not
    # enough headroom for another call -- proving Continue sees the SAME
    # accumulated total rather than a fresh per-operation budget.
    tightened = service.update_cost_ceiling(updated.job_id, manual_spend + Decimal("0.001"))

    done = service.continue_llm_generation(tightened.job_id, provider_factory=dispatch_factory(Dispatch()))
    assert done.status == "cost_ceiling"
    assert Decimal(done.accumulated_cost_usd) == manual_spend
    pending_after = [s for s in done.slots if s.kind == "llm" and s.status != "accepted"]
    assert len(pending_after) == 2  # neither originally-planned slot could run


# --------------------------------------------------------------------------- #
# 13. interim numbering is compact; final numbering is canonical and stable
# --------------------------------------------------------------------------- #
def test_interim_numbering_is_compact_final_numbering_is_canonical_and_stable(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {
                "מבוא": {"total": 2, "database": 2, "llm": 0},           # canonical idx 0
                "גרעיני הבסיס": {"total": 2, "database": 0, "llm": 2},    # canonical idx 6 -- all pending
                "היסטולוגיה": {"total": 1, "database": 1, "llm": 0},     # canonical idx 7
            }, "cost_ceiling_usd": "5.00", "_seed": 12,
        })
    view = service.result_view(job)
    assert [q["number"] for q in view["questions"]] == [1, 2, 3]
    assert [q["category"] for q in view["questions"]] == ["מבוא", "מבוא", "היסטולוגיה"]

    # the persisted FINAL number already reserves room for the pending
    # גרעיני הבסיס LLM slots in canonical-order position (3, 4) even though
    # they are invisible to the interim (compact) view above
    histology_slot = next(s for s in job.slots if s.category == "היסטולוגיה")
    assert histology_slot.number == 5

    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    final_view = service.result_view(done)
    assert [q["number"] for q in final_view["questions"]] == [1, 2, 3, 4, 5]
    assert [q["category"] for q in final_view["questions"]] == [
        "מבוא", "מבוא", "גרעיני הבסיס", "גרעיני הבסיס", "היסטולוגיה",
    ]


# --------------------------------------------------------------------------- #
# 14. DOCX/full-Excel/LLM-only-Excel behave correctly in both phases
# --------------------------------------------------------------------------- #
def test_exports_exclude_pending_placeholders_then_include_everything_after_continue(
    jobs_app, jobs_root, llm_ready,
):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 3, "database": 1, "llm": 2}},
            "cost_ceiling_usd": "5.00", "_seed": 13,
        })
    # full export: only the 1 accepted DB question, compact interim number
    full_ws = load_workbook(io.BytesIO(service.export_full_xlsx(job))).active
    rows = list(full_ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][0] == 1

    # LLM-only export: header row only, nothing generated yet
    llm_ws = load_workbook(io.BytesIO(service.export_llm_xlsx(job))).active
    assert llm_ws.max_row == 1

    # result_view()["questions"] is exactly what DOCX export receives
    # (frontend strips job-only metadata but never adds/removes rows)
    assert len(service.result_view(job)["questions"]) == 1

    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))

    full_ws2 = load_workbook(io.BytesIO(service.export_full_xlsx(done))).active
    rows2 = list(full_ws2.iter_rows(min_row=2, values_only=True))
    assert len(rows2) == 3
    assert [r[0] for r in rows2] == [1, 2, 3]  # final stable numbers

    llm_ws2 = load_workbook(io.BytesIO(service.export_llm_xlsx(done))).active
    assert llm_ws2.max_row == 3  # header + 2 llm rows

    assert len(service.result_view(done)["questions"]) == 3


# --------------------------------------------------------------------------- #
# 15. reload/restart restores db_review without starting generation
# --------------------------------------------------------------------------- #
def test_reload_restart_restores_db_review_without_starting_generation(jobs_app, jobs_root):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 14,
        })
    reloaded = store.load(job.job_id)
    assert reloaded.status == "queued"
    assert service.workflow_phase(reloaded) == "db_review"
    assert reloaded.cost_ledger == []

    # simulate a backend restart -- a db_review job was never "running", so
    # recover_on_start must leave it completely untouched
    changed = store.recover_on_start()
    assert job.job_id not in changed
    still = store.load(job.job_id)
    assert still.status == "queued"
    assert service.workflow_phase(still) == "db_review"
    assert still.cost_ledger == []


# --------------------------------------------------------------------------- #
# 16. awaiting/running jobs cannot be branched; completed jobs still can
# --------------------------------------------------------------------------- #
def test_branching_blocked_until_llm_phase_completes(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 15,
        })
    assert service.workflow_phase(job) == "db_review"
    with pytest.raises(service.JobConflict):
        service.branch_job(job.job_id, None)

    running = store.load(job.job_id)
    running.status = "running"
    store.save(running)
    assert service.workflow_phase(running) == "llm_generation"
    with pytest.raises(service.JobConflict):
        service.branch_job(job.job_id, None)

    interrupted = store.load(job.job_id)
    interrupted.status = "interrupted"
    store.save(interrupted)
    assert service.workflow_phase(interrupted) == "llm_generation"
    with pytest.raises(service.JobConflict):
        service.branch_job(job.job_id, None)

    # resume from "interrupted" via Continue, reaching genuine completion
    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    assert done.status == "completed"
    child = service.branch_job(done.job_id, {"mode": "custom", "custom_name": "בדיקת ענף"})
    assert child.status == "completed"
    assert child.parent_job_id == done.job_id


# --------------------------------------------------------------------------- #
# 17. legacy persisted jobs without the new phase load safely
# --------------------------------------------------------------------------- #
def test_legacy_job_json_without_workflow_phase_field_loads_and_maps_safely(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
            "cost_ceiling_usd": "5.00", "_seed": 16,
        })
    done = service.continue_llm_generation(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    raw_path = store._job_file(done.job_id)
    raw_before = raw_path.read_text("utf-8")
    assert "workflow_phase" not in raw_before  # never persisted -- pure function of status
    reloaded = store.load(done.job_id)
    assert service.workflow_phase(reloaded) == "complete"
    assert raw_path.read_text("utf-8") == raw_before  # reading never rewrites the file

    # a legacy job "stuck" queued with a still-queued LLM slot (a pre-WP27
    # crash between create_job's own save and run_job's first save) now reads
    # as db_review -- a strict improvement (it gets a working Continue
    # affordance it never had before), not a regression.
    with jobs_app.app_context():
        stuck = service.create_job({
            "categories": {CAT: {"total": 1, "database": 0, "llm": 1}},
            "cost_ceiling_usd": "5.00",
        })
    assert stuck.status == "queued"
    assert service.workflow_phase(stuck) == "db_review"

    # the other rare legacy corner: "queued" with NO llm slot at all reads as
    # complete, never stuck "awaiting approval" forever for zero planned items
    synthetic = store.load(stuck.job_id)
    synthetic.slots = [s for s in synthetic.slots if s.kind != "llm"]
    assert service.workflow_phase(synthetic) == "complete"


# --------------------------------------------------------------------------- #
# 18. exclusions / global DB matching compose correctly with db_review
# --------------------------------------------------------------------------- #
def test_exclusions_compose_correctly_with_db_review(jobs_app, jobs_root):
    with jobs_app.app_context():
        from src.models.question import Question

        rows = Question.query.filter(Question.category == CAT).all()
        excluded = [rows[0].id]
    with jobs_app.app_context():
        job = service.create_job({
            "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
            "cost_ceiling_usd": "5.00", "excluded_db_ids": excluded,
        })
    assert service.workflow_phase(job) == "db_review"
    selected = {s.db_id for s in job.slots if s.kind == "database"}
    assert selected.isdisjoint(set(excluded))
    assert job.excluded_db_ids == sorted(excluded)

    db_slot = next(s for s in job.slots if s.kind == "database")
    with jobs_app.app_context():
        updated = service.replace_from_db(
            job.job_id, db_slot.instance_id, extra_exclude_ids=(db_slot.db_id,),
        )
    new_slot = updated.slot_by_instance(db_slot.instance_id)
    assert new_slot.db_id not in excluded
    assert new_slot.db_id != db_slot.db_id


# --------------------------------------------------------------------------- #
# route-level: POST /exam-jobs/<id>/continue-llm
# --------------------------------------------------------------------------- #
def test_route_continue_llm_happy_path_and_conflicts(jobs_client, jobs_app, jobs_root, llm_ready):
    jobs_app.config["EXAM_JOB_SYNC"] = True
    jobs_app.config["EXAM_JOB_PROVIDER_FACTORY"] = dispatch_factory(Dispatch())
    resp = jobs_client.post("/api/exam-jobs", json={
        "categories": {CAT: {"total": 2, "database": 1, "llm": 1}},
        "cost_ceiling_usd": "5.00", "_seed": 17,
    })
    assert resp.status_code == 202
    body = resp.get_json()
    assert body["status"] == "queued"
    assert body["workflow_phase"] == "db_review"
    jid = body["job_id"]

    view = jobs_client.get(f"/api/exam-jobs/{jid}").get_json()
    assert view["workflow_phase"] == "db_review"
    assert [q["origin"] for q in view["questions"]] == ["database"]

    bad = jobs_client.post("/api/exam-jobs/11111111-1111-1111-1111-111111111111/continue-llm")
    assert bad.status_code == 404

    cont = jobs_client.post(f"/api/exam-jobs/{jid}/continue-llm")
    assert cont.status_code == 202, cont.get_json()

    view2 = jobs_client.get(f"/api/exam-jobs/{jid}").get_json()
    assert view2["workflow_phase"] == "complete"
    assert view2["status"] == "completed"
    assert sorted(q["origin"] for q in view2["questions"]) == ["database", "llm"]

    again = jobs_client.post(f"/api/exam-jobs/{jid}/continue-llm")
    assert again.status_code == 409


def test_route_b_zero_returns_completed_with_no_workflow_phase_db_review(jobs_client, jobs_app, jobs_root):
    jobs_app.config["EXAM_JOB_SYNC"] = True
    resp = jobs_client.post("/api/exam-jobs", json={
        "categories": {CAT: {"total": 2, "database": 2, "llm": 0}},
        "cost_ceiling_usd": "5.00", "_seed": 18,
    })
    assert resp.status_code == 202
    body = resp.get_json()
    assert body["status"] == "completed"
    assert body["workflow_phase"] == "complete"

    cont = jobs_client.post(f"/api/exam-jobs/{body['job_id']}/continue-llm")
    assert cont.status_code == 409
