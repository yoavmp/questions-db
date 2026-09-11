"""WP21 -- restored per-question/overall analytics, complete Excel exports,
ledger-derived attempt counts, and audit-overwrite prevention.

Offline: temp DB, fake in-process providers, sockets blocked (conftest). Zero
network / provider calls.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from tests._wp18_fakes import Dispatch, RejectingProvider, dispatch_factory

from src.jobs import service, store
from src.jobs.model import SEVEN, missing_analytics
from src.models.question import Question
from src.models.user import db

CAT = "היסטולוגיה"


def _seed_accuracy(jobs_app, category, accuracy_list, distinction_list=None):
    """Same history on every row in `category` -- deterministic regardless of
    which one random selection later picks."""
    with jobs_app.app_context():
        rows = Question.query.filter(Question.category == category).all()
        for r in rows:
            r.accuracy_list = list(accuracy_list)
            r.distinction_list = list(distinction_list or [])
        db.session.commit()


def _create(jobs_app, **cats):
    with jobs_app.app_context():
        return service.create_job({"categories": cats, "cost_ceiling_usd": "5.00"})


def _mixed(jobs_app, disp=None):
    disp = disp or Dispatch()
    with jobs_app.app_context():
        job = service.create_job(
            {"categories": {CAT: {"total": 4, "database": 2, "llm": 2}},
             "cost_ceiling_usd": "5.00", "_seed": 3}
        )
    done = service.run_job(job.job_id, provider_factory=dispatch_factory(disp))
    assert done.status == "completed"
    return done


# --------------------------------------------------------------------------- #
# §3 -- per-question analytics: zero / one / multiple historic values, LLM missing
# --------------------------------------------------------------------------- #
def test_db_selection_with_no_history_is_missing_not_fabricated(jobs_app, jobs_root):
    job = _create(jobs_app, **{CAT: {"total": 1, "database": 1, "llm": 0}})
    slot = next(s for s in job.slots if s.kind == "database")
    assert slot.analytics == missing_analytics()
    with jobs_app.app_context():
        dto = service.result_view(job)["questions"][0]
    assert dto["accuracy"] is None and dto["accuracy_list"] == []
    assert dto["distinction"] is None and dto["distinction_list"] == []


def test_db_selection_with_one_historic_value(jobs_app, jobs_root):
    _seed_accuracy(jobs_app, CAT, [85.0], [0.4])
    job = _create(jobs_app, **{CAT: {"total": 1, "database": 1, "llm": 0}})
    slot = next(s for s in job.slots if s.kind == "database")
    assert slot.analytics["accuracy"] == 85.0
    assert slot.analytics["accuracy_list"] == [85.0]
    assert slot.analytics["distinction"] == 0.4
    assert slot.analytics["distinction_list"] == [0.4]


def test_db_selection_with_multiple_historic_values(jobs_app, jobs_root):
    _seed_accuracy(jobs_app, CAT, [70.0, 90.0], [0.2, 0.6])
    job = _create(jobs_app, **{CAT: {"total": 1, "database": 1, "llm": 0}})
    slot = next(s for s in job.slots if s.kind == "database")
    assert slot.analytics["accuracy_list"] == [70.0, 90.0]
    assert slot.analytics["accuracy"] == 80.0     # mean, matches Question.accuracy
    assert slot.analytics["distinction_list"] == [0.2, 0.6]


def test_llm_question_always_missing_analytics(jobs_app, jobs_root, llm_ready):
    _seed_accuracy(jobs_app, CAT, [99.0])   # DB history exists but is irrelevant to the LLM slot
    with jobs_app.app_context():
        job = service.create_job(
            {"categories": {CAT: {"total": 1, "database": 0, "llm": 1}}, "cost_ceiling_usd": "5.00"}
        )
    done = service.run_job(job.job_id, provider_factory=dispatch_factory(Dispatch()))
    slot = next(s for s in done.slots if s.kind == "llm")
    assert slot.analytics == missing_analytics()
    with jobs_app.app_context():
        dto = service.result_view(done)["questions"][0]
    assert dto["accuracy"] is None and dto["accuracy_list"] == []


def test_analytics_persists_through_a_store_round_trip(jobs_app, jobs_root):
    _seed_accuracy(jobs_app, CAT, [88.0])
    job = _create(jobs_app, **{CAT: {"total": 1, "database": 1, "llm": 0}})
    reloaded = store.load(job.job_id)
    slot = next(s for s in reloaded.slots if s.kind == "database")
    assert slot.analytics["accuracy"] == 88.0
    assert slot.analytics["accuracy_list"] == [88.0]


# --------------------------------------------------------------------------- #
# §3 -- the four replacement transitions
# --------------------------------------------------------------------------- #
def test_replace_db_to_db_uses_the_new_rows_own_analytics(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        for i, r in enumerate(Question.query.filter(Question.category == CAT).all()):
            r.accuracy_list = [float(10 * (i + 1))]  # distinct per row
        db.session.commit()
    done = _mixed(jobs_app)
    db_slot = next(s for s in done.slots if s.kind == "database")
    with jobs_app.app_context():
        updated = service.replace_from_db(done.job_id, db_slot.instance_id)
    new_slot = updated.slot_by_instance(db_slot.instance_id)
    assert new_slot.kind == "database"
    assert new_slot.db_id != db_slot.db_id
    with jobs_app.app_context():
        row = Question.query.get(new_slot.db_id)
        expected = service._db_analytics(row)
    assert new_slot.analytics == expected
    assert new_slot.analytics["accuracy"] is not None


def test_replace_db_to_llm_analytics_becomes_missing(jobs_app, jobs_root, llm_ready):
    _seed_accuracy(jobs_app, CAT, [77.0])
    done = _mixed(jobs_app)
    db_slot = next(s for s in done.slots if s.kind == "database")
    assert db_slot.analytics["accuracy"] == 77.0
    updated = service.replace_via_llm(
        done.job_id, db_slot.instance_id, provider_factory=dispatch_factory(Dispatch()),
    )
    new_slot = updated.slot_by_instance(db_slot.instance_id)
    assert new_slot.kind == "llm"
    assert new_slot.analytics == missing_analytics()


def test_replace_llm_to_llm_analytics_stays_missing(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    llm_slot = next(s for s in done.slots if s.kind == "llm" and s.status == "accepted")
    assert llm_slot.analytics == missing_analytics()
    updated = service.replace_via_llm(
        done.job_id, llm_slot.instance_id, provider_factory=dispatch_factory(Dispatch()),
    )
    new_slot = updated.slot_by_instance(llm_slot.instance_id)
    assert new_slot.kind == "llm"
    assert new_slot.analytics == missing_analytics()


def test_replace_llm_to_db_uses_the_selected_rows_analytics(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        for i, r in enumerate(Question.query.filter(Question.category == CAT).all()):
            r.accuracy_list = [float(5 * (i + 1))]
        db.session.commit()
    done = _mixed(jobs_app)
    llm_slot = next(s for s in done.slots if s.kind == "llm" and s.status == "accepted")
    with jobs_app.app_context():
        updated = service.replace_from_db(done.job_id, llm_slot.instance_id)
    new_slot = updated.slot_by_instance(llm_slot.instance_id)
    assert new_slot.kind == "database"
    with jobs_app.app_context():
        row = Question.query.get(new_slot.db_id)
        expected = service._db_analytics(row)
    assert new_slot.analytics == expected
    assert new_slot.analytics["accuracy"] is not None


def test_failed_replace_llm_leaves_analytics_unchanged(jobs_app, jobs_root, llm_ready):
    _seed_accuracy(jobs_app, CAT, [42.0])
    done = _mixed(jobs_app)
    db_slot = next(s for s in done.slots if s.kind == "database")
    before = dict(db_slot.analytics)
    updated = service.replace_via_llm(
        done.job_id, db_slot.instance_id,
        provider_factory=dispatch_factory(Dispatch(factory=RejectingProvider)),
    )
    new_slot = updated.slot_by_instance(db_slot.instance_id)
    assert new_slot.kind == "database"     # unchanged
    assert new_slot.analytics == before    # unchanged, not reset to missing


def test_analytics_never_reaches_generator_context(jobs_app, jobs_root, llm_ready):
    _seed_accuracy(jobs_app, CAT, [55.0])
    done = _mixed(jobs_app)
    llm_slot = next(s for s in done.slots if s.kind == "llm" and s.status == "accepted")
    prev = service._previous_for_slot(done, llm_slot)
    assert prev, "expected at least the DB question(s) as previous-question context"
    for p in prev:
        assert set(p.keys()) == set(SEVEN)
        assert "analytics" not in p and "accuracy" not in p and "id" not in p


# --------------------------------------------------------------------------- #
# §5 -- two Excel exports
# --------------------------------------------------------------------------- #
def _rows(xlsx_bytes: bytes) -> list:
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    return list(ws.iter_rows(min_row=1, values_only=True))


def test_full_export_schema_order_and_missing_value_convention(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        seeded_row = Question.query.filter(Question.category == CAT).first()
        seeded_row.accuracy_list = [80.0, 90.0]
        seeded_row.distinction_list = [0.4]
        seeded_id = seeded_row.id
        db.session.commit()
    done = _mixed(jobs_app)
    rows = _rows(service.export_full_xlsx(done))
    assert rows[0] == tuple(service.FULL_EXPORT_HEADERS)

    ordered_slots = done.accepted_questions_ordered()
    body = rows[1:]
    assert len(body) == len(ordered_slots) == 4          # 2 DB + 2 LLM, current, canonical order
    assert [r[0] for r in body] == [1, 2, 3, 4]           # מספר_שאלה, canonical global numbers

    # openpyxl round-trips a written "" back as None -- verified empirically;
    # either way there is no fabricated id/measurement, just a blank cell.
    for r, slot in zip(body, ordered_slots):
        if slot.kind == "llm":
            assert not r[1]    # מזהה_שאלה missing, never fabricated
            assert not r[9]    # דיוק missing
            assert not r[10]   # הבחנה missing
        else:
            assert isinstance(r[1], int) and r[1] == slot.db_id
            if slot.db_id == seeded_id:
                assert r[9] == "[80.0, 90.0]"   # exact legacy multi-value convention
                assert r[10] == "[0.4]"
            else:
                assert not r[9] and not r[10]  # this DB row has no history either


def test_full_export_never_inserts_into_the_db(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        before = Question.query.count()
    done = _mixed(jobs_app)
    service.export_full_xlsx(done)
    with jobs_app.app_context():
        assert Question.query.count() == before


def test_llm_only_export_unchanged_by_wp21(jobs_app, jobs_root, llm_ready):
    """Regression guard: the WP18 upload-compatible export keeps its schema
    and row set exactly (WP21 introduces export_full_xlsx alongside it, not
    instead of it)."""
    done = _mixed(jobs_app)
    rows = _rows(service.export_llm_xlsx(done))
    assert rows[0] == tuple(service.EXPORT_HEADERS)
    llm_count = sum(1 for s in done.accepted_questions_ordered() if s.kind == "llm")
    assert len(rows) - 1 == llm_count


# --------------------------------------------------------------------------- #
# §7 -- audit-overwrite prevention
# --------------------------------------------------------------------------- #
def test_repeated_operations_on_one_slot_never_overwrite_earlier_evidence(jobs_app, jobs_root, llm_ready):
    with jobs_app.app_context():
        job = service.create_job(
            {"categories": {CAT: {"total": 1, "database": 0, "llm": 1}}, "cost_ceiling_usd": "5.00"}
        )
    # 1st invocation: forced rejection
    done = service.run_job(job.job_id, provider_factory=dispatch_factory(Dispatch(factory=RejectingProvider)))
    slot = next(s for s in done.slots if s.kind == "llm")
    assert slot.status == "failed"
    first_ref = slot.audit_ref
    first_dir = store.job_dir(done.job_id) / first_ref
    assert first_dir.is_dir()
    first_manifest_before = (first_dir / "attempt_01" / "manifest.json").read_bytes()

    # 2nd invocation on the SAME slot: a successful retry
    retried = service.retry_slot(done.job_id, slot.slot_id, provider_factory=dispatch_factory(Dispatch()))
    slot2 = retried.slot_by_id(slot.slot_id)
    assert slot2.status == "accepted"
    second_ref = slot2.audit_ref
    second_dir = store.job_dir(retried.job_id) / second_ref

    assert second_ref != first_ref                      # distinct invocation directories
    assert first_dir.is_dir() and second_dir.is_dir()    # both still exist
    # the first invocation's on-disk evidence is byte-for-byte untouched
    assert (first_dir / "attempt_01" / "manifest.json").read_bytes() == first_manifest_before

    # both invocations are linked to distinct, resolvable ledger records
    ledger_refs = [e["audit_ref"] for e in retried.cost_ledger]
    assert first_ref in ledger_refs and second_ref in ledger_refs
    assert len(set(ledger_refs)) == len(ledger_refs)     # every invocation ref is unique
    for ref in (first_ref, second_ref):
        assert ".." not in ref and not ref.startswith("/")   # safe relative reference only


def test_replace_llm_also_gets_its_own_invocation_directory(jobs_app, jobs_root, llm_ready):
    done = _mixed(jobs_app)
    llm_slot = next(s for s in done.slots if s.kind == "llm" and s.status == "accepted")
    initial_ref = llm_slot.audit_ref
    updated = service.replace_via_llm(
        done.job_id, llm_slot.instance_id, provider_factory=dispatch_factory(Dispatch()),
    )
    new_slot = updated.slot_by_instance(llm_slot.instance_id)
    assert new_slot.audit_ref != initial_ref
    initial_dir = store.job_dir(done.job_id) / initial_ref
    assert initial_dir.is_dir()  # the original accepted call's own evidence untouched
